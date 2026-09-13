"""End-to-end CLI entry point.

    python -m src.run_analysis

Loads the raw CSV, cleans it, writes the cleaned dataset + audit log and a
self-contained HTML dashboard, then prints a short business summary.
"""

from __future__ import annotations

import json

import pandas as pd

from . import (
    analysis,
    config,
    dashboard,
    data_cleaning,
    data_loader,
    feature_engineering,
    ml,
    visualization,
)


def _write_workbook(df: pd.DataFrame, log_rows: list, path) -> None:
    """Write cleaned data + cleaning log, with light Excel styling."""
    log_df = pd.DataFrame(log_rows)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Cleaned_Data", index=False)
        log_df.to_excel(writer, sheet_name="Cleaning_Log", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(path)
    zebra = PatternFill("solid", fgColor="F2F2F2")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = zebra
        for col in range(1, ws.max_column + 1):
            letter = ws.cell(row=1, column=col).column_letter
            longest = max(
                (
                    len(str(ws.cell(row=r, column=col).value or ""))
                    for r in range(1, ws.max_row + 1)
                ),
                default=8,
            )
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), 45)
        ws.freeze_panes = "A2"
    wb.save(path)


def _section(title: str) -> None:
    print("\n" + "=" * 78)
    print(title)
    print("=" * 78)


def main() -> None:
    _section("1. LOAD & CLEAN")
    raw = data_loader.load_and_validate()
    print(f"Raw file: {config.RAW_DATA_FILE}  shape={raw.shape}")
    clean, log, metrics = data_cleaning.clean_orders(raw, verbose=True)

    clean.to_csv(config.CLEANED_CSV, index=False)
    pd.DataFrame(log).to_csv(config.CLEANING_LOG_FILE, index=False)
    _write_workbook(clean, log, config.CLEANED_XLSX)
    print(f"Cleaned CSV  : {config.CLEANED_CSV}")
    print(f"Cleaned XLSX : {config.CLEANED_XLSX}")
    print(f"Cleaning log : {config.CLEANING_LOG_FILE} ({len(log)} actions)")

    _section("2. FEATURE ENGINEERING")
    df = feature_engineering.add_features(clean)
    summary = feature_engineering.summarise_features(df)
    print(f"Date range : {summary['date_min']} -> {summary['date_max']}")
    print(f"Years      : {summary['years']}")
    print(f"Products   : {summary['products']}   Countries: {summary['countries']}")

    _section("3. ANALYSIS")
    kpis = analysis.overview_metrics(df)
    by_product = analysis.revenue_by_product(df)
    trend = analysis.revenue_trend(df)
    seasonality = analysis.monthly_seasonality(df)
    by_country = analysis.revenue_by_country(df)
    by_payment = analysis.revenue_by_payment(df)
    status_df = analysis.status_breakdown(df)
    customers = analysis.top_customers(df)
    impact = analysis.cancellation_impact(df)
    by_note = analysis.revenue_by_note(df)

    print(f"Total orders        : {kpis['total_orders']:,}")
    print(f"Recognised orders   : {kpis['recognised_orders']:,} ({kpis['completion_rate']:.0f}%)")
    print(f"Total revenue       : ${kpis['total_revenue']:,.2f}")
    print(f"Average order value : ${kpis['avg_order_value']:,.2f}")
    print(f"Units sold          : {kpis['total_units']:,}")
    print(f"Cancelled / Refunded: {kpis['cancelled_orders']} / {kpis['refunded_orders']}")
    print("\nTop 5 products by revenue:")
    print(by_product.head(5).to_string(index=False))
    print("\nRevenue by country:")
    print(by_country.to_string(index=False))

    _section("4. MACHINE LEARNING")

    print("Algorithm 1 - order-outcome classifier (predicts recognised revenue):")
    ml_result = ml.train_order_outcome_model(df)
    print(ml_result["cv_table"].to_string(index=False))
    model_metrics = ml_result["metrics"]
    print(f"\nBest challenger : {model_metrics['best_model']}")
    print(
        f"Held-out test   : accuracy={model_metrics['accuracy']:.2f}  "
        f"ROC AUC={model_metrics['roc_auc']:.2f}  "
        f"precision={model_metrics['precision']:.2f}  "
        f"recall={model_metrics['recall']:.2f}  f1={model_metrics['f1']:.2f} "
        f"(train={model_metrics['n_train']}, test={model_metrics['n_test']})"
    )
    if model_metrics["beats_baseline"]:
        conclusion = "weakly predictable from order attributes."
    else:
        conclusion = (
            "not clearly predictable from order attributes alone - the model does "
            "not beat a naive baseline, which is itself a finding."
        )
    print(
        f"Cross-validated AUC = {model_metrics['cv_roc_auc']:.2f} vs baseline "
        f"{model_metrics['baseline_auc']:.2f} -> order outcome is {conclusion}"
    )
    print("\nTop predictors:")
    print(ml_result["feature_importance"].head(6).to_string(index=False))

    print("\nAlgorithm 2 - RFM customer segmentation (KMeans):")
    seg_result = ml.segment_customers(df)
    print(
        f"Clusters chosen by silhouette: k={seg_result['chosen_k']} "
        f"(score={seg_result['silhouette']})"
    )
    print(seg_result["summary"].to_string(index=False))

    config.ML_METRICS_FILE.write_text(
        json.dumps(
            {
                "order_outcome": model_metrics,
                "cv_table": ml_result["cv_table"].to_dict(orient="records"),
                "segmentation": {
                    "chosen_k": seg_result["chosen_k"],
                    "silhouette": seg_result["silhouette"],
                },
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    _section("5. CHARTS & DASHBOARD")
    champion = seg_result["summary"].iloc[0]
    hero = visualization.chart_hero(
        {
            "title": "E-commerce Orders Lens",
            "subtitle": "Cleaning  ·  EDA  ·  Machine Learning  ·  Dashboard",
            "metrics": [
                (
                    f"{metrics['rows_after']} clean orders",
                    f"from {metrics['rows_before']} raw rows",
                ),
                (f"${kpis['total_revenue']:,.0f}", "recognised revenue"),
                (
                    f"{champion['revenue_share_%']:.0f}% revenue",
                    f"from {champion['customers']} Champions",
                ),
                ("13 charts", "one HTML dashboard"),
            ],
        }
    )
    chart_paths = [
        visualization.chart_top_bottom_products(by_product),
        visualization.chart_revenue_trend(trend),
        visualization.chart_revenue_by_country(by_country),
        visualization.chart_revenue_by_payment(by_payment),
        visualization.chart_status_breakdown(status_df),
        visualization.chart_top_customers(customers),
        visualization.chart_seasonality(seasonality),
        visualization.chart_cancellation_impact(impact),
        visualization.chart_quality_before_after(metrics),
        visualization.chart_ml_confusion_matrix(
            ml_result["confusion"], model_metrics["best_model"]
        ),
        visualization.chart_ml_feature_importance(ml_result["feature_importance"]),
        visualization.chart_ml_roc(ml_result["roc"], model_metrics["best_model"]),
        visualization.chart_customer_segments(seg_result["summary"]),
    ]
    for path in chart_paths:
        print(f"  chart -> {path.name}")
    print(f"  banner -> {hero.name}")

    tables = {
        "products": by_product,
        "countries": by_country,
        "payments": by_payment,
        "customers": customers,
        "status": status_df,
        "segments": seg_result["summary"].rename(
            columns={"segment": "customer_segment", "customers": "customers"}
        ),
        "predictors": ml_result["feature_importance"].head(10),
    }
    dashboard_file = dashboard.build_dashboard(metrics, kpis, tables, chart_paths)
    print(f"Dashboard    : {dashboard_file}")
    print(f"ML metrics   : {config.ML_METRICS_FILE}")

    _section("6. EXTRA INSIGHTS")
    print("Revenue by payment method:")
    print(by_payment.to_string(index=False))
    print("\nValue locked in non-revenue statuses:")
    print(impact.to_string(index=False))
    print("\nAdd-on impact:")
    print(by_note.to_string(index=False))
    print("\nDone.")


if __name__ == "__main__":
    main()
